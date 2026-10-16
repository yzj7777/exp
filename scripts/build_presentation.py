from io import BytesIO
from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1]
PDF=ROOT/'docs/reference/openJiuwen-创意大赛提交模板-参考.pdf'
OUT=ROOT/'deliverables/需求翻译官-比赛项目介绍.pptx'
# Match the landscape-A4 proportions of the supplied PDF reference.
SW,SH=11.69,8.27
FONT='Microsoft YaHei'
C={'white':'FFFFFF','ink':'171717','red':'C90000','blue':'1683D8','muted':'565656','line':'85857C','purple':'7664E8','lav':'F1EEFF','mint':'42A983','mintbg':'EAF7F1','gold':'D59A3D','goldbg':'FFF7E8','soft':'F7F6FA','navy':'292A40'}
prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH); blank=prs.slide_layouts[6]

def rgb(v):return RGBColor.from_string(v)
def text(slide,s,x,y,w,h,size=12,color=None,bold=False,italic=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=.02):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin);tf.margin_top=tf.margin_bottom=Inches(margin);tf.vertical_anchor=valign
    p=tf.paragraphs[0];p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
    r=p.add_run();r.text=s;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.italic=italic;r.font.color.rgb=rgb(color or C['ink'])
    return box
def rect(slide,x,y,w,h,fill=None,line=None,rounded=False,transparency=0):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:
        sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.fill.transparency=transparency
    else: sh.fill.background()
    if line: sh.line.color.rgb=rgb(line);sh.line.width=Pt(.7)
    else: sh.line.fill.background()
    if rounded:
        try:sh.adjustments[0]=.1
        except Exception:pass
    return sh
def line(slide,x1,y1,x2,y2,color='B8B8B3',width=1,dash=False):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    if dash:
        try:sh.line.dash_style=2
        except Exception:pass
    return sh
def add_footer(slide,n,logo):
    text(slide,str(n),.65,7.93,.4,.15,7,'3F3F3F')
    if logo: slide.shapes.add_picture(BytesIO(logo), Inches(9.93), Inches(7.77), width=Inches(1.22), height=Inches(.31))
def title(slide,s): text(slide,s,.34,.82,10.95,.51,25,C['red'],True)
def tag_text(slide,label,required=False,x=.55,y=1.7,w=3):
    text(slide,label,x,y,w,.3,14,C['ink'])
    if required: text(slide,'（必填）',x+len(label)*.20,y,w=.95,h=.3,size=14,color=C['red'],bold=True)

def add_cover(doc,logo):
    slide=prs.slides.add_slide(blank)
    p=doc[0];pix=p.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False)
    slide.shapes.add_picture(BytesIO(pix.tobytes('png')),0,0,width=Inches(SW),height=Inches(SH))
    # Cover the template placeholder title cleanly, then fill with the project title.
    rect(slide,.20,3.52,6.0,.8,C['white'])
    text(slide,'需求翻译官',.35,3.55,5.3,.55,27,C['ink'],True)
    text(slide,'openJiuwen 智能体创新大赛 · 创意方案',.38,4.15,5.5,.28,12,C['muted'])
    text(slide,'把模糊想法，翻译成可确认、可讨论的需求。',.38,4.55,6.3,.35,14,C['blue'],False,True)
    return slide

def add_slide2(logo):
    s=prs.slides.add_slide(blank);title(s,'创意名称：需求翻译官')
    # template body panel
    rect(s,.08,1.55,11.52,5.35,None,C['line'],True)
    line(s,5.58,1.58,5.58,6.88,'B6B6B1',.8,True)
    line(s,.09,5.77,5.58,5.77,'B6B6B1',.8,True)
    tag_text(s,'创意描述',True,.42,1.83,2.3)
    text(s,'（500字以内）',2.05,1.84,1.45,.28,10,C['muted'])
    desc=('需求翻译官是一款面向产品需求沟通的交互式 Agent。它通过少量、循序渐进的问答，帮助需求提出者把模糊想法逐步澄清为结构化需求草稿。Agent 聚焦目标用户、痛点、预期结果、关键流程、范围约束和验收方式；回答含糊时继续追问，暂时无法确定的信息标为待确认。最终由用户核对、修改后再交给产品或研发团队评审。项目坚持“先澄清，再整理；不确定就标注，不替用户做决定”，降低需求误解与反复确认的风险。')
    text(s,desc,.48,2.35,4.8,3.12,11,C['blue'],False,True)
    tag_text(s,'所属主题',True,.48,5.98,2.3)
    text(s,'办公',.7,6.38,1.2,.28,14,C['blue'],False,True)
    tag_text(s,'应用场景',True,5.95,1.83,2.2)
    scenario=('面向业务部门提出软件功能需求、产品经理整理需求的沟通场景。业务方往往先说“做得更方便”或“加一个审批功能”，却未明确使用者、触发条件、主流程、边界和成功标准。需求翻译官通过交互访谈逐项澄清，形成可复核的需求摘要、功能范围、未决问题与验收草稿。潜在价值是让业务、产品与研发围绕同一份有来源、有状态的需求材料继续协作，避免把 Agent 推测误当成用户已确认的事实。目标用户与实际成效仍需通过真实访谈验证。')
    text(s,scenario,5.95,2.55,5.2,3.65,12,C['blue'],False,True)
    add_footer(s,2,logo);return s

def add_slide3(logo):
    s=prs.slides.add_slide(blank);title(s,'创意方案构想')
    rect(s,.08,1.55,11.52,5.35,None,C['line'],True)
    line(s,.09,4.63,11.58,4.63,'B6B6B1',.8,True);line(s,5.58,4.64,5.58,6.88,'B6B6B1',.8,True)
    tag_text(s,'方案详情',True,.55,1.82,2.3)
    body=('核心能力：\n'
          '1. 接收自然语言需求，围绕用户、痛点、目标、流程、范围和验收方式建立需求画像。\n'
          '2. 优先询问会影响需求理解的关键缺口；每轮少量提问，避免一次性长问卷。\n'
          '3. 发现笼统回答时追问具体例子；允许跳过，并把未知信息标为“待确认”。\n'
          '4. 记录回答来源与确认状态，阶段性总结并请用户核对、修改。\n'
          '5. 生成结构化需求草稿，供需求方确认、产品经理评审后再进入后续流程。')
    text(s,body,.55,2.28,10.5,2.12,12,C['blue'],False,True)
    tag_text(s,'方案优势（可选）',False,.4,4.85,2.7)
    advantages=('与直接扩写一句提示词不同，本方案把“澄清过程”作为核心：问题围绕需求缺口推进；用户事实、待确认项和 Agent 建议分开记录；草稿可追溯、可修订。当前已有可交互规则原型与完整角色模拟案例，但真实用户效果尚待评估。')
    text(s,advantages,.4,5.42,4.85,1.18,10,C['blue'],False,True)
    tag_text(s,'技术依赖（可选）',False,5.9,4.85,2.7)
    deps=('当前：原生 HTML/CSS/JavaScript 规则驱动原型，无模型 API、后端或 WorkSwarm 集成。\n后续：若入围，计划评估服务端 LLM 接入、结构化会话状态、人工确认与隐私保护；模型服务和预算待定。现有样例测试为手写规则检查，不代表真实效果。')
    text(s,deps,5.9,5.42,5.1,1.18,10,C['blue'],False,True)
    add_footer(s,3,logo);return s

def add_slide4(logo):
    s=prs.slides.add_slide(blank)
    title(s,'DEMO 或示意图：一次需求澄清如何发生')
    text(s,'示意案例：用户说“想让团队会议协作更高效”——以下是访谈路径设计，模拟回答不代表真实用户调研。',.45,1.42,10.7,.35,10,C['muted'])
    # Four-step UI journey
    xs=[.38,3.24,6.1,8.96]; w=2.34; y=2.15; h=3.4
    titles=['01 说出想法','02 Agent 追问','03 标记信息状态','04 生成需求草稿']
    bodies=[
      ('输入：','“我想让会议更高效。”\n\n用户不需要先会写 PRD。'),
      ('问题：','“哪个环节最耗时或最容易遗漏？谁在使用？”\n\n一次聚焦少量关键问题。'),
      ('画像：','已记录：会议协作\n待确认：具体痛点\n建议：举一个最近的例子'),
      ('输出：','用户与场景\n痛点与目标\n主流程与范围\n未决项与验收标准')]
    for i in range(4):
        rect(s,xs[i],y,w,h,'FBFBFC','B8B8B1',True)
        text(s,titles[i],xs[i]+.17,y+.22,w-.34,.35,12,C['red'],True)
        line(s,xs[i]+.14,y+.72,xs[i]+w-.14,y+.72,'D6D6D2',.8)
        for j,(lab,val) in enumerate([bodies[i][0:2]]):
            text(s,lab,xs[i]+.17,y+1.03,w-.34,.23,10,C['ink'],True)
            text(s,val,xs[i]+.17,y+1.35,w-.34,1.65,10,C['blue'],False,True)
        if i<3:
            line(s,xs[i]+w+.06,y+1.7,xs[i+1]-.1,y+1.7,C['purple'],1.7)
            # arrow head
            tri=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(xs[i+1]-.18),Inches(y+1.57),Inches(.17),Inches(.27));tri.fill.solid();tri.fill.fore_color.rgb=rgb(C['purple']);tri.line.fill.background()
    rect(s,.42,5.92,10.82,.62,C['goldbg'],None,True)
    text(s,'当前可操作 Demo：web/showcase.html → 现场输入 → 追问 → 查看并下载 Markdown 草稿。当前网页为规则驱动原型，未调用 LLM API。',.62,6.07,10.4,.3,9,'8C682D',True)
    add_footer(s,4,logo);return s

def add_reference_slide(doc,index,logo=None):
    s=prs.slides.add_slide(blank);p=doc[index];pix=p.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False)
    s.shapes.add_picture(BytesIO(pix.tobytes('png')),0,0,width=Inches(SW),height=Inches(SH))
    return s

def build():
    if not PDF.exists():raise FileNotFoundError(f'Missing supplied template reference: {PDF}')
    source=pymupdf.open(PDF)
    logo=None
    for img in source[1].get_images(full=True):
        data=source.extract_image(img[0])
        if data.get('width',0)>100 and data.get('height',0)<100: logo=data['image'];break
    add_cover(source,logo)
    add_slide2(logo);add_slide3(logo);add_slide4(logo)
    # Keep the official reference appendix pages (赛道说明、评分标准、Thank you) as supplied.
    add_reference_slide(source,4);add_reference_slide(source,5);add_reference_slide(source,6)
    OUT.parent.mkdir(parents=True,exist_ok=True);prs.save(OUT);print(OUT)
if __name__=='__main__':build()
